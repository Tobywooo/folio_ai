"""Text extraction for PDF, PPTX, TXT, and web pages."""

from __future__ import annotations

import io
import ipaddress
import socket
from dataclasses import dataclass
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from .errors import IngestionError, ValidationError

_MAX_PAGE_BYTES = 5 * 1024 * 1024
_MAX_REDIRECTS = 5
_USER_AGENT = "FolioAI/1.0 (+notebook source ingestion)"


@dataclass
class Section:
    """A unit of extracted text. `page` is the PDF page or PPTX slide number (1-based), else None."""

    text: str
    page: int | None = None


def extract_pdf(data: bytes) -> list[Section]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestionError("This PDF is password-protected.")
        return [Section(page.extract_text() or "", number) for number, page in enumerate(reader.pages, 1)]
    except IngestionError:
        raise
    except Exception as exc:  # pypdf raises a variety of errors on malformed files
        raise IngestionError(f"Could not read PDF: {exc}") from exc


def extract_pptx(data: bytes) -> list[Section]:
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"Could not read PowerPoint file: {exc}") from exc

    sections = []
    for number, slide in enumerate(deck.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        sections.append(Section("\n".join(p for p in parts if p.strip()), number))
    return sections


def _shape_text(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [text for child in shape.shapes for text in _shape_text(child)]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [shape.text_frame.text]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    return []


def extract_txt(data: bytes) -> list[Section]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return [Section(data.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [Section(data.decode("latin-1", errors="replace"))]


# --------------------------------------------------------------------------- web


@dataclass
class WebPage:
    url: str  # final URL after redirects
    title: str
    html: bytes
    sections: list[Section]


def fetch_url(url: str) -> WebPage:
    """Download a page and extract its readable text.

    Redirects are followed manually so every hop is checked against private/internal
    addresses (otherwise a URL could point the Function App at internal Azure endpoints).
    """
    current = url.strip()
    for _ in range(_MAX_REDIRECTS + 1):
        _check_public_url(current)
        try:
            resp = requests.get(
                current, headers={"User-Agent": _USER_AGENT}, timeout=20, stream=True, allow_redirects=False
            )
        except requests.RequestException as exc:
            raise IngestionError(f"Could not fetch URL: {exc}") from exc
        if resp.is_redirect:
            location = resp.headers.get("Location", "")
            resp.close()
            if not location:
                raise IngestionError("URL redirected without a destination.")
            current = urljoin(current, location)
            continue
        break
    else:
        raise IngestionError("URL redirected too many times.")

    with resp:
        if resp.status_code >= 400:
            raise IngestionError(f"URL returned HTTP {resp.status_code}.")
        content_type = resp.headers.get("Content-Type", "").lower()
        if content_type and not any(t in content_type for t in ("text/html", "text/plain", "application/xhtml")):
            raise IngestionError(
                f"URL content type '{content_type.split(';')[0]}' isn't supported. "
                "For PDFs or slides, download the file and upload it instead."
            )
        body = bytearray()
        for block in resp.iter_content(64 * 1024):
            body.extend(block)
            if len(body) > _MAX_PAGE_BYTES:
                raise IngestionError("Web page is larger than 5 MB.")

    html = bytes(body)
    if "text/plain" in content_type:
        return WebPage(current, current, html, extract_txt(html))

    soup = BeautifulSoup(html, "html.parser")
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "aside", "form", "iframe"]):
        tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.body or soup
    lines = [line.strip() for line in root.get_text("\n").splitlines()]
    text = "\n".join(line for line in lines if line)
    return WebPage(current, title or current, html, [Section(text)])


def _check_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValidationError("Enter a full http:// or https:// URL.")
    try:
        infos = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise IngestionError(f"Could not resolve host '{parsed.hostname}'.") from exc
    for info in infos:
        address = ipaddress.ip_address(info[4][0].split("%")[0])
        if not address.is_global:
            raise ValidationError("URLs pointing at private or internal network addresses aren't allowed.")
